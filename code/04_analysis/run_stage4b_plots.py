"""
Stage 4b — re-run plot + metrics only (UMAP already done, reload npy).
"""
from __future__ import annotations
import sys, json, time, gc
from pathlib import Path

import numpy as np
import pandas as pd
import torch
import matplotlib.pyplot as plt
from sklearn.metrics import silhouette_score
from sklearn.neighbors import NearestNeighbors
from scipy.stats import spearmanr

ROOT = Path('/DATA/lxk/lxkzero/claude/qd/data')
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT/'result'))
from _style import TIMES, NC_PALETTE, SEQ_CMAP, neat_axes, PERSIST_C  # noqa
import _style  # noqa

OUT  = ROOT/'result/stage4_latent'
FIGS = ROOT/'result/figs'
OUT.mkdir(parents=True, exist_ok=True)
FIGS.mkdir(parents=True, exist_ok=True)
LOG  = ROOT/'result/logs/stage4.log'

def log(msg, end='\n'):
    s = f'[{time.strftime("%H:%M:%S")}] {msg}'
    print(s, end=end, flush=True)
    with open(LOG, 'a') as f: f.write(s + end)

# ====== Load UMAP coords ======
log('loading saved UMAP / latent')
U1 = np.load(OUT/'umap_E1.npy')
U0 = np.load(OUT/'umap_E0.npy')
E1 = np.load(OUT/'latent_E1.npy')
N = len(U1)
log(f'  U1={U1.shape}, U0={U0.shape}, E1={E1.shape}')

# ====== Rebuild metadata (30000-row aligned) ======
log('rebuilding metadata')
preds = pd.read_csv(ROOT/'model_data/model1_out/model1_predictions.csv')
preds = preds[preds['smiles_valid'] == True].reset_index(drop=True)
master = pd.read_csv(ROOT/'03_transport/PFAS_master_structure.csv', low_memory=False)
strict_smi = set(master.loc[master['pfas_class'].notna(), 'SMILES'].dropna().astype(str).str.strip())
preds['__smi__'] = preds['SMILES'].astype(str).str.strip()
preds['is_strict'] = preds['__smi__'].isin(strict_smi)

strict_idx = preds.index[preds['is_strict']].tolist()
other_idx  = preds.index[~preds['is_strict']].tolist()
rng = np.random.default_rng(42)
N_SAMPLE = 30000
keep_other = rng.choice(other_idx, size=min(N_SAMPLE - len(strict_idx), len(other_idx)), replace=False)
keep = sorted(set(strict_idx) | set(keep_other.tolist()))
sub = preds.loc[keep].reset_index(drop=True)
assert len(sub) == N, f'{len(sub)} != {N}'

meta_df = pd.DataFrame({'SMILES': sub['SMILES'].values})
_preds_slim = preds[['SMILES','pred_logKow','pred_water_solubility','is_strict']].drop_duplicates('SMILES')
meta_df = meta_df.merge(_preds_slim, on='SMILES', how='left')
m2_pred = pd.read_csv(ROOT/'model_data/model2_out/model2_predictions.csv',
                      usecols=['SMILES','pred_logBAF','pred_logBCF','pred_logBCFD'])
m2_pred = m2_pred.drop_duplicates('SMILES')
meta_df = meta_df.merge(m2_pred, on='SMILES', how='left')
meta_df = meta_df.iloc[:N].reset_index(drop=True)

for c in ['pred_logBCF','pred_logBAF','pred_logKow','pred_water_solubility']:
    if c in meta_df:
        s = meta_df[c]
        q1, q99 = s.quantile([0.01, 0.99])
        meta_df[c+'_clip'] = s.clip(q1, q99)

# ClassyFire
cf = pd.read_csv(ROOT/'01_structure/classyfire/classyfire_PFAS_pool_subset.csv')
if 'chem_cf_subclass_name' in cf.columns:
    cf = cf.rename(columns={'chem_cf_subclass_name':'cf_subclass', 'chem_cf_class_name':'cf_class'})
mast_ik = master[['SMILES','InChIKey']].dropna().drop_duplicates('SMILES')
cf2 = cf.merge(mast_ik, on='InChIKey', how='left').dropna(subset=['SMILES'])
cf_slim = cf2[['SMILES','cf_subclass','cf_class']].drop_duplicates('SMILES')
meta_df = meta_df.merge(cf_slim, on='SMILES', how='left')
meta_df = meta_df.iloc[:N].reset_index(drop=True)

# Persistence
pers = pd.read_csv(ROOT/'05_persistence_transformation/persistence.csv')
comp = pd.read_csv(ROOT/'05_persistence_transformation/compounds.csv', on_bad_lines='skip')
pers_merge = pers.merge(comp[['compound_id','cas']], on='compound_id', how='left')
pers_by_cas = pers_merge.groupby('cas')['persistence_label'].first().reset_index()
mast2 = master[['SMILES','CAS']].dropna().copy()
mast2['cas'] = mast2['CAS'].astype(str).str.upper().str.strip()
pers_by_cas['cas'] = pers_by_cas['cas'].astype(str).str.upper().str.strip()
pers_smi = mast2.merge(pers_by_cas, on='cas', how='inner')[['SMILES','persistence_label']].drop_duplicates('SMILES')
meta_df = meta_df.merge(pers_smi, on='SMILES', how='left')
meta_df = meta_df.iloc[:N].reset_index(drop=True)

meta_df.to_csv(OUT/'meta.csv', index=False)
log(f'  meta OK  shape={meta_df.shape}  persistence={meta_df["persistence_label"].notna().sum()}  cf={meta_df["cf_subclass"].notna().sum()}')
assert len(meta_df) == N

# ====== Plotting helpers ======
def scatter_continuous(U, vals, title, cbar_label, fname, cmap=SEQ_CMAP):
    fig, ax = plt.subplots(figsize=(7.5, 6.0))
    mask = np.isfinite(vals)
    if mask.sum() < 10:
        plt.close(); return
    sc = ax.scatter(U[mask, 0], U[mask, 1], c=vals[mask], s=4, alpha=0.7,
                    cmap=cmap, edgecolors='none', rasterized=True)
    cb = fig.colorbar(sc, ax=ax, fraction=0.045, pad=0.03)
    cb.set_label(cbar_label, fontsize=13)
    ax.set_xlabel('UMAP dim 1', labelpad=6); ax.set_ylabel('UMAP dim 2', labelpad=6)
    ax.set_title(title, pad=16)
    neat_axes(ax)
    plt.tight_layout(pad=2.0)
    plt.savefig(FIGS/fname, dpi=300); plt.savefig(OUT/fname, dpi=300)
    plt.close()

log('plotting fig 4a (logKow)')
scatter_continuous(U1, meta_df['pred_logKow_clip'].values,
    'Transport latent space – predicted log $K_{ow}$',
    'predicted log $K_{ow}$', 'fig4a_logKow.png')

log('plotting fig 4b (logBCF)')
scatter_continuous(U1, meta_df['pred_logBCF_clip'].values,
    'Transport latent space – predicted log BCF',
    'predicted log BCF', 'fig4b_logBCF.png')

log('plotting fig 4c (persistence)')
fig, ax = plt.subplots(figsize=(7.5, 6.0))
labels = meta_df['persistence_label'].fillna('unlabelled').values
nan_mask = (labels == 'unlabelled')
ax.scatter(U1[nan_mask, 0], U1[nan_mask, 1], c='#E0E0E0', s=2.5, alpha=0.4,
           edgecolors='none', rasterized=True, label='unlabelled')
for i, lab in enumerate(['vP','P','not-P_self','not-P','NP']):
    m = (labels == lab)
    if not m.any(): continue
    ax.scatter(U1[m, 0], U1[m, 1], c=PERSIST_C.get(lab, NC_PALETTE[i]),
               s=22, alpha=0.95, edgecolors='white', linewidths=0.4, label=lab)
ax.set_xlabel('UMAP dim 1', labelpad=6); ax.set_ylabel('UMAP dim 2', labelpad=6)
ax.set_title('Transport latent space – persistence label', pad=16)
ax.legend(frameon=False, loc='upper right', fontsize=11, borderpad=0.5)
neat_axes(ax)
plt.tight_layout(pad=2.0)
plt.savefig(FIGS/'fig4c_persistence.png', dpi=300); plt.savefig(OUT/'fig4c_persistence.png', dpi=300)
plt.close()
log('  fig 4c done')

log('plotting fig 4d (ClassyFire subclass)')
sub_mask = meta_df['is_strict'].values & meta_df['cf_subclass'].notna().values
fig, ax = plt.subplots(figsize=(9.5, 6.0))
ax.scatter(U1[~sub_mask, 0], U1[~sub_mask, 1], c='#E0E0E0', s=2.5, alpha=0.35,
           edgecolors='none', rasterized=True, zorder=1)
top = pd.Series(meta_df.loc[sub_mask, 'cf_subclass']).value_counts().head(8).index.tolist()
for i, lab in enumerate(top):
    m = (meta_df['cf_subclass'] == lab).values
    ax.scatter(U1[m, 0], U1[m, 1], c=NC_PALETTE[i % len(NC_PALETTE)], s=18,
               alpha=0.92, edgecolors='white', linewidths=0.3, label=str(lab), zorder=2+i)
other = sub_mask & (~meta_df['cf_subclass'].isin(top).values)
ax.scatter(U1[other, 0], U1[other, 1], c='#9E9E9E', s=8, alpha=0.55,
           edgecolors='none', label='other strict PFAS', zorder=2)
ax.set_xlabel('UMAP dim 1', labelpad=6); ax.set_ylabel('UMAP dim 2', labelpad=6)
ax.set_title('Transport latent space – ClassyFire subclass (strict PFAS)', pad=16)
ax.legend(loc='center left', bbox_to_anchor=(1.02, 0.5), frameon=False,
          fontsize=9.5, handlelength=1.2, borderpad=0.5)
neat_axes(ax)
plt.tight_layout(pad=2.0)
plt.savefig(FIGS/'fig4d_subclass.png', dpi=300); plt.savefig(OUT/'fig4d_subclass.png', dpi=300)
plt.close()
log('  fig 4d done')

# ====== Quantitative metrics ======
log('computing quantitative metrics E0 vs E1')

def latent_distance_corr(E, target, n_pairs=8000, seed=0):
    rng = np.random.default_rng(seed)
    finite = np.where(np.isfinite(target))[0]
    if len(finite) < 50: return np.nan
    pairs = rng.choice(finite, size=(n_pairs, 2))
    keep = pairs[:, 0] != pairs[:, 1]
    pairs = pairs[keep]
    dE = np.linalg.norm(E[pairs[:,0]] - E[pairs[:,1]], axis=1)
    dY = np.abs(target[pairs[:,0]] - target[pairs[:,1]])
    rho, _ = spearmanr(dE, dY)
    return rho

def silhouette_by(E, labels):
    s = pd.Series(labels)
    keep = s.notna() & (s.astype(str).str.lower() != 'unlabelled') & (s.astype(str) != 'nan')
    s = s[keep]; Ek = E[keep.values]
    counts = s.value_counts()
    big = counts[counts >= 8].index.tolist()
    if len(big) < 2: return np.nan
    keep2 = s.isin(big)
    if keep2.sum() < 20: return np.nan
    if keep2.sum() > 8000:
        rng = np.random.default_rng(0)
        idx = rng.choice(np.where(keep2)[0], size=8000, replace=False)
        return float(silhouette_score(Ek[idx], s.iloc[idx].values))
    return float(silhouette_score(Ek[keep2.values], s[keep2].values))

def knn_consistency(E, labels, k=15):
    s = pd.Series(labels)
    keep = s.notna() & (s.astype(str).str.lower() != 'unlabelled') & (s.astype(str) != 'nan')
    if keep.sum() < 50: return np.nan
    Ek = E[keep.values]
    sk = s[keep].values
    nbrs = NearestNeighbors(n_neighbors=k+1).fit(Ek)
    _, ind = nbrs.kneighbors(Ek)
    same = (sk[ind[:,1:]] == sk[:,None]).mean()
    return float(same)

# E0: rebuild PCA
from sklearn.decomposition import PCA
from sklearn.preprocessing import StandardScaler
import featurize as _ft
import openpyxl

log('  re-featurize for E0')
meta_json = json.load(open(ROOT/'model_data/model1_out/model1_meta.json'))
in_mean_m1 = np.asarray(meta_json['in_mean'])
in_std_m1  = np.asarray(meta_json['in_std'])
fill_med_m1 = np.asarray(meta_json['fill_median'])

wb = openpyxl.load_workbook(ROOT/'model_data/Model1_Transport.xlsx', data_only=True)
fstats_sheet = wb['F_stats_dict']
fstats_h = [c.value for c in next(fstats_sheet.iter_rows(min_row=1, max_row=1))]
fstats_r = list(fstats_sheet.iter_rows(min_row=2, values_only=True))
fstats_df = pd.DataFrame(fstats_r, columns=fstats_h)
ftz = _ft.Featurizer()
X_raw, valid_mask = ftz.transform(sub['SMILES'].tolist(), fstats_df=fstats_df, fstats_key='SMILES')
# for E0 metrics use the valid subset matching U1
# Since U1 was already built on a sub that was fully valid (valid = 30000), we can use X_raw directly
Xp = X_raw.copy()
nm = ~np.isfinite(Xp)
Xp[nm] = np.broadcast_to(fill_med_m1, Xp.shape)[nm]
Xp = (Xp - in_mean_m1) / np.where(in_std_m1 > 1e-8, in_std_m1, 1.0)
pca = PCA(n_components=64, random_state=42)
E0 = pca.fit_transform(Xp)
log(f'  E0 shape={E0.shape}  explained_var={pca.explained_variance_ratio_[:5].sum():.3f}')

rows = []
for E, name in [(E0, 'E0 structure-only (PCA-64)'), (E1, 'E1 transport latent (128d)')]:
    for c in ['pred_logKow', 'pred_logBCF', 'pred_logBAF']:
        rho = latent_distance_corr(E, meta_df[c].values)
        rows.append({'embedding':name, 'metric':f'Spearman(d_E, |Δ {c}|)', 'value':round(rho,4) if not np.isnan(rho) else np.nan})
    rows.append({'embedding':name, 'metric':'silhouette by ClassyFire subclass',
                 'value': round(silhouette_by(E, meta_df['cf_subclass'].values), 4)})
    rows.append({'embedding':name, 'metric':'silhouette by persistence label',
                 'value': round(silhouette_by(E, meta_df['persistence_label'].values), 4)})
    rows.append({'embedding':name, 'metric':'kNN-15 consistency by subclass',
                 'value': round(knn_consistency(E, meta_df['cf_subclass'].values), 4)})
    rows.append({'embedding':name, 'metric':'kNN-15 consistency by persistence',
                 'value': round(knn_consistency(E, meta_df['persistence_label'].values), 4)})

qm = pd.DataFrame(rows)
qm.to_csv(OUT/'quant_metrics.csv', index=False)
log('Quantitative metrics:')
log(qm.to_string(index=False))

log('Stage 4 (plot+metrics) done.')
