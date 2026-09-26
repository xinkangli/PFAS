from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt,Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import zipfile,json,hashlib
W=Path(__file__).resolve().parent;BASE=W.parents[1]/'outputs/wr_core';O=W.parents[1]/'outputs/wr_v17';O.mkdir(exist_ok=True)
def replace(z,text):
 rpr=deepcopy(z.runs[0]._r.rPr) if z.runs and z.runs[0]._r.rPr is not None else None
 for e in list(z._p):
  if e.tag!=qn('w:pPr'):z._p.remove(e)
 r=z.add_run(text)
 if rpr is not None:r._r.insert(0,rpr)
def save(d,path,ref):
 d.save(path)
 with zipfile.ZipFile(path) as z:parts={n:z.read(n) for n in z.namelist()}
 with zipfile.ZipFile(ref) as z:
  for n in z.namelist():
   if n in ['word/styles.xml','word/numbering.xml'] or n.startswith(('word/header','word/footer')):parts[n]=z.read(n)
 with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
  for n,b in parts.items():z.writestr(n,b)
ref=BASE/'paper_WR_V16_en.docx';d=Document(ref)
changes=[]
for z in d.paragraphs:
 t=z.text
 if t.startswith('The historical modeling ledger contains'):
  replace(z,'The original compiled database contains 2,404 observations from ECOTOX and ITRC. Replaying the archived screening rules retained 2,343 records with nonempty, parsable structures, 1,925 with a confirmed positive value convertible to L/kg, 1,209 with wet-weight basis, and 1,078 after the initial censoring-note screen. These sequential counts describe data eligibility, not independent sample sizes. The resulting 1,078-row ledger contains 23 compounds and 51 source DOIs; all entries in this screened subset originate from the ITRC aquatic compilation. Re-execution of the original ITRC parser reproduced every ledger value from the saved Excel cells, and all 1,078 identifiers and historical values matched both the SQLite database and the modeling workbook. Database-level traceability is therefore complete for this subset, whereas verification against the original experimental literature is a separate requirement.')
  changes.append('2.1 collection and screening pathway')
 elif t.startswith('The reconstruction preserves source locators'):
  replace(z,'Original-paper numerical locators were established for 382 ledger entries. Of these, 216 met the frozen endpoint and value-interpretation criteria, while 166 were excluded because of unresolved censoring, analytical or whole-body reconstruction, plasma-unit incompatibility, or other documented value/basis issues. The remaining 696 records from 43 source DOIs are traceable to database cells but have not been verified against exact original-paper values: 689 lack a completed original-table check, four have unresolved aggregation, and three have unresolved numerical mapping. They are excluded from all current primary training and evaluation partitions. The eligible population contains 147 BAF and 69 BCF observations for 19 compounds from six studies (Table 1; Figure 1). Exclusion is based on evidence status and endpoint compatibility, not prediction error; the excluded records remain in a separate audit archive.')
  changes.append('2.1 frozen admission and exclusion')
 elif t.startswith('A targeted review corrected 80'):
  replace(z,t+' The 69 retained BCF targets are tissue-specific steady-state logBCF values verified in Chen et al. (2016), Supplementary Table S9; the original log values are used at their reported precision. They are not arithmetic range midpoints. The 69 range-derived midpoint records in the historical ledger are a different set, and all are excluded from the primary analysis. No label, partition, or model prediction changed during the present provenance closure.')
  changes.append('2.1 correction of BCF provenance description')
 elif t.startswith('Source reconstruction changes the effective scale'):
  replace(z,'Source reconstruction distinguishes a traceable compiled record from a source-qualified model target. All 1,078 screened ledger values can be reproduced from the archived ITRC workbook, but only 216 satisfy the current primary criteria after original-source checks. The 696 records awaiting exact original-table verification are not missing-source records and are not used as training labels. Another 166 source-located entries remain excluded for specified endpoint or value-interpretation reasons. This closes the inclusion decision for every ledger row without implying that the historical literature audit is complete. The 216 eligible observations represent only 19 compounds, with all primary BCF observations arising from one study. Repeated measurements characterize observed conditions but do not establish broad independent chemical or study coverage.')
  changes.append('3.1 traceability versus experimental verification')
 elif t.startswith('The central limitations are the small number'):
  replace(z,t.replace('incomplete verification of the historical ledger','limited original-paper verification beyond the admitted primary set'))
 elif t.startswith('The revision archive contains'):
  replace(z,'The V17 source-closure archive contains a primary-only 216-row dataset, a separate 862-row exclusion archive, a row-level disposition ledger, original-code replay checks, source hashes, and the sequential selection counts. Its directory is data/nc/old/water_research_v17_source_closure. The unchanged model results, frozen partitions, weights, prediction-level outputs and figure sources remain in data/nc/old/water_research_core_20260925. Admission checks confirmed that none of the excluded records entered the 28 current primary/similarity jobs or three calibration splits. Historical raw files are preserved. This server working release retains upstream dependencies and is not described as an already deposited self-contained public repository.')
# Chen primary source citation at the paragraph that uses it, without perturbing existing numbering.
for z in d.paragraphs:
 if 'Chen et al. (2016), Supplementary Table S9;' in z.text:replace(z,z.text.replace('Chen et al. (2016), Supplementary Table S9;','Chen et al. (2016), Supplementary Table S9 [13];'))
# Copy complete verified Chen bibliography entry from the reference used for V16.
v15=Document(W.parents[1]/'outputs/v15_validation/paper_v15_validated_en.docx')
chen=next(z for z in v15.paragraphs if z.text.startswith('29.') and 'Chen' in z.text)
p=deepcopy(chen._p);d._element.body.insert(-1,p);replace(d.paragraphs[-1],'13. '+chen.text.split('.',1)[1].strip())
save(d,O/'paper_WR_V17_en.docx',ref)
ref=BASE/'support_WR_V16_en.docx';si=Document(ref)
for z in si.paragraphs:
 t=z.text
 if t.startswith('Version WR V16'):replace(z,t.replace('Version WR V16','Version WR V17').replace('current context-corrected evaluation','current context-corrected evaluation and completed row-level admission audit'))
 elif t.startswith('The primary population is selected'):
  replace(z,'The primary population consists of the 216 records admitted by the frozen v15_primary flag; a primary-only file is now supplied as data/primary_216_only.csv in the V17 closure archive. The original final_observations.csv remains an audit ledger and is not an instruction to train on every row. All 1,078 screened records have reproducible ITRC database-cell lineage. Exact original-paper numerical locators were established for 382 records; 216 were admitted and 166 were excluded for stated interpretation or endpoint reasons. The other 696 records are database-traced but not exact-original-table verified. None was used in the current primary analysis. Section S7 provides the complete selection and disposition accounting.')
 elif t.startswith('The primary BCF panel contains 69'):
  replace(z,'The primary BCF panel comprises 69 tissue-specific steady-state logBCF values from Chen et al. (2016), DOI 10.1016/j.scitotenv.2016.05.215, Supplementary Table S9. The original log values are retained at the reported precision, with tissue and exposure group preserved. All 69 current targets match the saved original-table extraction; all originated from directly reported database values rather than range midpoint conversion. By contrast, the historical ledger contains a separate set of 69 arithmetic range midpoints, all excluded from the primary population. The V16 description conflated these two sets; that wording is corrected here without changing model labels or results. Plasma ratios, reconstructed whole-body estimates, censored summaries and unresolved records remain excluded under their documented reasons.')
 elif t.startswith('The server training environment uses'):
  replace(z,t+' The V17 admission audit adds no biological labels and does not retrain an unchanged analysis. Its primary-only export and exclusion archive are separate from the V16 feature matrix, whose row ordering is retained; they must be joined by record_id rather than substituted positionally.')
# Append native template paragraphs for the new audit section.
proto=next(z for z in si.paragraphs if z.text.startswith('The primary population'))
head=next(z for z in si.paragraphs if z.text.startswith('S1.'))
def add(text,heading=False):
 e=deepcopy((head if heading else proto)._p);si._element.body.insert(-1,e);z=si.paragraphs[-1];replace(z,text);return z
add('S7. Original-code replay and closed admission decisions',True)
add('The original collector disi/collectors/itrc.py downloads the ITRC aquatic workbook and parses its BCF-BAF Database sheet. Endpoint values, source citation, species/tissue fields and database UID are then passed through disi/schema.py into pfas_bcf_master. data/build_model_data.py joins structure information by normalized CAS and exports X_Y_joined. The archived audit_data.py screening rules generate the 1,078-row provisional ledger. The present closure script replays the parser without writing to the original database and independently reconstructs numeric cell values, including the documented range rule. It then checks each record against the saved SQLite value and modeling workbook by identifier. Every one of the 1,078 records passes all four checks. This verifies the retained extraction pathway, not the entire historical collection of 2,404 records or the execution time of an undocumented original search.')
add('The parser can substitute an arithmetic midpoint when only minimum and maximum values are present; such records must not be represented as directly measured point observations. It also contains generic qualifier-stripping and unit-default logic. The current audit preserves raw cells and conversion rules instead of treating those defaults as experimental confirmation. Primary-table and endpoint qualification remain independent gates. All 69 midpoint-derived ledger rows are outside the current primary population.')
def table(headers,rows,widths):
 t=si.add_table(rows=1,cols=len(headers));t.autofit=False
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 for c,v in zip(t.rows[0].cells,headers):c.text=v
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 for i,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
  for c in row.cells:
   for z in c.paragraphs:
    z.alignment=0;z.paragraph_format.line_spacing=1;z.paragraph_format.space_after=Pt(4)
    for r in z.runs:r.font.name='Times New Roman';r.font.size=Pt(9);r.bold=i==0
add('Table S7. Sequential reconstruction of the screening population. Exclusions are calculated between consecutive stages.')
flow=[['Compiled modeling workbook',2404,'—'],['Nonempty and parsable structure',2343,61],['Positive value with confirmed L/kg conversion',1925,418],['Wet-weight basis',1209,716],['No unresolved censoring note in initial screen',1078,131],['Source-qualified primary population',216,862]]
table(['Stage','Retained','Excluded at stage'],flow,[3.75,.8,1.2])
add('Table S8. Final disposition of all 1,078 screened records. Source location alone does not establish eligibility.')
table(['Disposition','Records','Current model use'],[['Source-qualified eligible observations',216,'Included'],['Database traced; original-table values not verified',696,'Excluded'],['Original-paper values located; endpoint or interpretation criteria not met',166,'Excluded'],['Total',1078,'216 included; 862 archived']],[3.45,.7,1.6])
add('The 696 database-traced exclusions comprise 689 records without a completed original-table numeric check, four with unresolved aggregation, and three with unresolved numeric mapping, spanning 43 source DOIs. The 166 source-located exclusions comprise 68 censored-summary observations, 46 reconstructed whole-body observations, 24 plasma-unit mismatches, 13 value/basis conflicts, seven LOQ-substituted summaries, six isomer-sum cases, and two alternative analytical estimates. These mutually exclusive categories sum to the exclusion total; they must not be interpreted as 862 fabricated or erroneous observations. Detailed locators and reasons appear in audit/all_1078_disposition.csv.')
add('Admission tests confirm zero overlap between the 862 excluded record identifiers and any train/test partition of the 28 primary/similarity jobs. The three calibration split files also contain only admitted identifiers. Current BCF targets were checked against all 69 saved original-table log values, with exact numerical agreement at the retained precision. The corrected primary-only export has the same values and identifiers as the population used for the V16 results. Consequently, no new performance estimate or significance claim is introduced by this closure.')
save(si,O/'support_WR_V17_en.docx',ref)
(O/'changes.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2))
print('V17_ENGLISH_DOCUMENTS_BUILT')
# Renumber main-paper citations in order of first appearance after adding Chen.
paper=O/'paper_WR_V17_en.docx';d=Document(paper);split=next(i for i,z in enumerate(d.paragraphs) if z.text=='References');paras=d.paragraphs
import re
refs={int(re.match(r'(\d+)\.',z.text).group(1)):z.text.split('.',1)[1].strip() for z in paras[split+1:] if re.match(r'\d+\.',z.text)};used=[]
def citation(m):
 nums=[int(x.strip()) for x in m.group(1).split(',')]
 for n in nums:
  assert n in refs
  if n not in used:used.append(n)
 return '['+', '.join(str(used.index(n)+1) for n in nums)+']'
for z in paras[:split]:
 if re.search(r'\[(\d+(?:,\s*\d+)*)\]',z.text):replace(z,re.sub(r'\[(\d+(?:,\s*\d+)*)\]',citation,z.text))
proto=deepcopy(paras[split+1]._p)
for z in paras[split+1:]:z._p.getparent().remove(z._p)
for i,n in enumerate(used,1):d._element.body.insert(-1,deepcopy(proto));replace(d.paragraphs[-1],f'{i}. {refs[n]}')
save(d,paper,BASE/'paper_WR_V16_en.docx')
p=O/'support_WR_V17_en.docx';si=Document(p)
for z in si.paragraphs:
 if 'The V16 description conflated' in z.text:replace(z,z.text.replace('The V16 description conflated these two sets; that wording is corrected here without changing model labels or results. ',''))
save(si,p,BASE/'support_WR_V16_en.docx')
# Concise Chinese decision record, kept separate from the submission manuscript.
ref=BASE/'WR_V16_核心修改说明_中文.docx';zh=Document(ref);bodyproto=deepcopy(zh.paragraphs[2]._p);headproto=deepcopy(zh.paragraphs[1]._p);titleproto=deepcopy(zh.paragraphs[0]._p)
for e in list(zh._element.body):
 if e.tag!=qn('w:sectPr'):zh._element.body.remove(e)
def zp(text,kind='body'):
 e=deepcopy({'body':bodyproto,'head':headproto,'title':titleproto}[kind]);zh._element.body.insert(-1,e);z=zh.paragraphs[-1];replace(z,text);z.paragraph_format.line_spacing=1.15;z.paragraph_format.space_after=Pt(6)
 for r in z.runs:
  r._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Serif CJK SC')
 for prop in z._p.xpath('./w:pPr/w:snapToGrid'):prop.getparent().remove(prop)
 off=OxmlElement('w:snapToGrid');off.set(qn('w:val'),'0');z._p.get_or_add_pPr().append(off)
zp('Water Research V17 数据溯源与纳入决定','title')
zp('结论','head')
zp('696 条记录有明确的 ITRC 数据库来源和可复现的原始代码路径，并非来源不明或伪造。缺的是逐条核实所引论文的原始数值。本轮已将它们正式列入排除清单，与另外 166 条虽定位到原文但不符合当前纳入标准的记录一起单独归档。主分析固定为 216 条，不删除任何原始文件，也不为增加样本量而降低核验标准。')
zp('一 代码追查结果','head')
zp('路径为 disi/collectors/itrc.py 读取 ITRC 水生 Excel → disi/schema.py 写入 SQLite → data/build_model_data.py 按 CAS 连接结构并形成 X_Y_joined → audit_data.py 执行单位、湿重和删失初筛 → 原文表格核对及端点资格审查。本轮重新运行原采集解析函数，并独立按 Excel 单元格重算；1,078 条全部与解析结果、数据库 ID/数值及建模工作表一致。')
zp('2,404 条原始编译记录中，2,343 条有可解析结构，1,925 条有可确认的正值及 L/kg 单位转换，1,209 条满足湿重口径，去除初筛中的删失说明后剩 1,078 条。这不是从 2,404 条直接跳到 216 条；完整逐阶段数量已写入支撑材料表 S7。采集代码完整回放确认的是这 1,078 条的数据库处理路径，不能据此宣称全部 2,404 条都已核实到原始实验。')
zp('二 每条记录如何处理','head')
zp('216 条：原文来源和当前端点/数值解释满足既定标准，作为唯一主分析表 primary_216_only.csv。包括 147 条 BAF 和 69 条 BCF，共 19 个化合物、6 项研究。')
zp('696 条：来自 43 个来源 DOI，全部保留数据库单元格和解析证据，但不纳入当前建模。其中 689 条未完成原表数值核对，4 条原表已找到但聚合方式无法重建，3 条原表已找到但数值映射未解决。清单为 excluded_696_database_traced.csv。排除不意味着原始测量错误；未来补足原文证据后可按同一规则重新审查。')
zp('166 条：虽已定位原文，但当前解释不合格，包括删失汇总 68 条、重建全身数值 46 条、血浆量纲不符 24 条、数值/基准冲突 13 条、LOQ 替代汇总 7 条、异构体求和 6 条、替代分析估计 2 条。它们单独列于 excluded_166_endpoint_or_value.csv。')
zp('三 是否影响已经重算的结果','head')
zp('不影响。本轮对 28 个主分析/相似性任务的训练和测试 ID 逐一检查，862 条排除记录的交集均为零；3 组校准划分也通过同样核对。216 条的标签、结构、记录身份、划分和 V16 模型结果均未改变，所以无需仅因归档重新训练。单独导出的主表不得按行号替换原特征矩阵，应按 record_id 对接。')
zp('需要纠正我上一版的一个错误：V16 把当前 69 条 BCF 写成范围中点，这是文字描述错误。它们实际来自 Chen 2016 原文 SI 表 S9 的组织稳态 logBCF，69 个当前标签均与已核实的原文值吻合。历史数据中另有 69 条范围中点，这另一组记录全部已排除。V17 正文与支撑材料已纠正这一点。')
zp('四 对形成 Water Research 投稿工作的意义','head')
zp('本轮完成的是数据纳入决策闭合：每条记录都有处理结果，不再把“全部旧数据逐篇追齐”设为无限期前置任务。数据库追溯与原始实验核验分层报告，既保留已有数据资产，也避免未核实标签进入核心证据。V17 保留 V16 的 5 幅主图和 2 张主表，支撑材料新增完整筛选及排除表。')
zp('这并不自动提高统计效力。主分析独立化合物仍为 19 个，BCF 仍集中于一项研究；运输监督的收益有限，跨研究区间仍跨零，外部误差仍需如实呈现。若要进一步增强投稿竞争力，优先补核能增加独立化合物或独立研究的原始数据，而不是为了记录数重新加入这 696 条。能否被 Water Research 接收仍取决于证据强度和编辑、审稿判断。')
zp('文件位置','head')
zp('/DATA/lxk/lxkzero/claude/qd/data/nc/old/water_research_v17_source_closure/。documents 为新版文稿；data 为纳入与隔离数据；audit 为逐记录清单和检查结果；code 保留原代码快照及本轮回放脚本。模型、权重和图源继续使用相邻的 water_research_core_20260925 目录，未重复计算或覆盖旧稿。')
save(zh,O/'WR_V17_数据溯源与处理说明_中文.docx',ref)
manifest={p.name:{'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'tables':len(Document(p).tables),'figures':len(Document(p).inline_shapes)} for p in O.glob('*.docx')}
assert manifest['paper_WR_V17_en.docx']['figures']==5
assert manifest['support_WR_V17_en.docx']['tables']==8
(O/'document_integrity.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print('V17_COMPLETE',list(manifest))
